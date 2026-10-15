#!/usr/bin/env python3
"""Convert PPTX slide pages to png cards and keep stable PowerPoint slide IDs."""
import json, subprocess, tempfile, shutil, sys
from pathlib import Path
from pptx import Presentation

root=Path(__file__).resolve().parents[1]
pptx=root/'cards'/'cards.pptx'
if not pptx.exists():
 print('No cards/cards.pptx found: leaving placeholder cards.json intact')
 sys.exit(0)
prs=Presentation(str(pptx))
assert len(prs.slides)>=1, 'PowerPoint must have at least one slide (one per card)'
overrides=json.loads((root/'card_overrides.json').read_text())
images=root/'card-images'
if images.exists():shutil.rmtree(images)
images.mkdir()
with tempfile.TemporaryDirectory() as tmp:
 tmp=Path(tmp)
 subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/lo-survey-convert','--headless','--convert-to','pdf','--outdir',str(tmp),str(pptx)],check=True)
 pdf=tmp/'cards.pdf'
 if not pdf.exists():raise RuntimeError('PowerPoint PDF conversion failed')
 subprocess.run(['pdftoppm','-png','-r','115',str(pdf),str(images/'page')],check=True)
 pages=sorted(images.glob('page-*.png'))
 assert len(pages)==len(prs.slides), f'Expected {len(prs.slides)} images, got {len(pages)}'
 cards=[]
 for i, slide in enumerate(prs.slides):
  card_id='SLIDE-'+str(prs.slides[i].slide_id)
  filename=f'{card_id}.png'
  pages[i].rename(images/filename)
  override=overrides.get(card_id,{})
  title=override.get('title',f'Card {i+1}')
  cards.append(dict(id=card_id,title=title,image='card-images/'+filename,active=override.get('active',True),options=override.get('options',['I would follow the instructions stated on the card','I would ignore the conditions printed on the card','I would skip this card entirely']),correct_option=override.get('correct_option',0)))
 (root/'cards.json').write_text(json.dumps(cards,ensure_ascii=False,indent=2)+'\n')
 print(f'Wrote {len(cards)} card images and cards.json')
